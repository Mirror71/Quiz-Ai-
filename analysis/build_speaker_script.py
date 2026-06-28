"""Builds Speaker_Script_Slides_1-2.docx — a spoken script the presenter can read.

Run:  python analysis/build_speaker_script.py
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "Speaker_Script_Slides_1-2.docx"

INDIGO = RGBColor(0x43, 0x38, 0xCA)
GRAY = RGBColor(0x55, 0x55, 0x55)

doc = Document()

# Base style
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(12)

# ---- Title
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("QuizAI — Presentation Speaker Script")
r.bold = True
r.font.size = Pt(20)
r.font.color.rgb = INDIGO

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
rs = sub.add_run("Your part: Slides 1 & 2  ·  Read naturally, don't rush")
rs.italic = True
rs.font.color.rgb = GRAY

doc.add_paragraph()


def heading(text, secs):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = INDIGO
    s = p.add_run(f"   (~{secs})")
    s.italic = True
    s.font.size = Pt(11)
    s.font.color.rgb = GRAY


def say(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.4
    p.add_run(text)


def tip(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    r = p.add_run("👉 Tip: " + text)
    r.italic = True
    r.font.size = Pt(11)
    r.font.color.rgb = GRAY


# ===================================================== SLIDE 1
heading("SLIDE 1 — Title & Introduction", "40–45 seconds")
tip("Smile, look at the audience, speak slowly for the first few seconds.")
say("“Good morning everyone. Thank you for being here. "
    "We are group [number / names], and today we are presenting our "
    "Artificial Intelligence project.”")
say("“Our project is called QuizAI. It is built under the topic of "
    "Generative AI.”")
say("“In one sentence: QuizAI is a web application that takes any PDF or "
    "PowerPoint document and automatically turns it into an interactive "
    "multiple-choice quiz, using Google’s Gemini AI model.”")
say("“So instead of writing practice questions by hand, you simply upload "
    "your study material, and the AI creates the quiz for you in a few "
    "seconds. The full project is also available on our GitHub repository.”")
say("“Let me start by explaining the problem we wanted to solve.”")
tip("That last line is your transition — click to Slide 2 as you say it.")

doc.add_paragraph()

# ===================================================== SLIDE 2
heading("SLIDE 2 — The Problem & Our Solution", "55–60 seconds")
tip("Point to the left side of the slide for the problem, the right side for the solution.")

p = doc.add_paragraph()
r = p.add_run("The Problem")
r.bold = True
r.font.color.rgb = INDIGO
say("“When students study from documents like lecture slides or textbooks, "
    "testing their own understanding is hard. Writing good quiz questions by "
    "hand is slow and takes effort — you need a correct answer, believable "
    "wrong options, and an explanation for each one.”")
say("“And for self-study, students need a fast and easy way to check what "
    "they actually understood after reading.”")

p = doc.add_paragraph()
r = p.add_run("Our Solution: QuizAI")
r.bold = True
r.font.color.rgb = INDIGO
say("“Our solution solves this with Generative AI. The user just uploads a "
    "PDF or PowerPoint file. QuizAI reads the content, and the Gemini model "
    "generates ten multiple-choice questions — each with four options, the "
    "correct answer, and a short explanation.”")
say("“Then the user takes the quiz directly in the app, one question at a "
    "time, and gets instant feedback showing whether they were right or "
    "wrong, plus a final score at the end.”")
say("“Now my teammate will explain how the system actually works behind the "
    "scenes.”")
tip("Final line hands over to the next presenter — turn slightly toward them.")

doc.add_paragraph()
foot = doc.add_paragraph()
rf = foot.add_run("Remember: breathe, pause between sentences, and it's okay to "
                  "glance at this sheet. You've got this! 🎯")
rf.italic = True
rf.font.color.rgb = GRAY

doc.save(OUT)
print("Wrote", OUT)
