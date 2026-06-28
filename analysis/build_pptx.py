"""Builds QuizAI_Presentation.pptx — clean white Office-theme look, 10 slides.

Run:  python analysis/build_pptx.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(__file__).resolve().parent.parent
FIG = ROOT / "analysis" / "figures"
OUT = ROOT / "QuizAI_Presentation.pptx"

ACCENT = RGBColor(0x44, 0x72, 0xC4)
ACCENT_DK = RGBColor(0x2E, 0x4E, 0x8A)
DARK = RGBColor(0x26, 0x26, 0x26)
GRAY = RGBColor(0x59, 0x59, 0x59)
CARD = RGBColor(0xF2, 0xF2, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x16, 0xA3, 0x4A)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height

TITLE_SLIDE = prs.slide_layouts[0]
TITLE_ONLY = prs.slide_layouts[5]


def rect(slide, x, y, w, h, color):
    shp = slide.shapes.add_shape(1, x, y, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = color
    shp.line.fill.background(); shp.shadow.inherit = False
    return shp


def textbox(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, (text, size, color, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(6)
        r = p.add_run(); r.text = text
        r.font.size = Pt(size); r.font.color.rgb = color
        r.font.bold = bold; r.font.name = "Calibri"
    return tb


def bullets(slide, x, y, w, h, items, size=20, color=DARK):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(11)
        r = p.add_run(); r.text = "•  " + item
        r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = "Calibri"
    return tb


def content_slide(title):
    slide = prs.slides.add_slide(TITLE_ONLY)
    th = slide.shapes.title
    th.left, th.top, th.width, th.height = Inches(0.6), Inches(0.3), Inches(12.1), Inches(1.0)
    p = th.text_frame.paragraphs[0]; p.text = title
    p.runs[0].font.size = Pt(32); p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = ACCENT_DK
    rect(slide, Inches(0.6), Inches(1.35), Inches(12.1), Pt(3), ACCENT)
    return slide


def flow(slide, y, items, boxh=Inches(0.85)):
    n = len(items); gap = Inches(0.55)
    boxw = Emu(int((SW - Inches(1.2) - gap * (n - 1)) / n))
    x = Inches(0.6)
    for i, it in enumerate(items):
        rect(slide, x, y, boxw, boxh, CARD)
        rect(slide, x, y, Inches(0.09), boxh, ACCENT)
        textbox(slide, x, y, boxw, boxh, [(it, 15, ACCENT_DK, True)],
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < n - 1:
            textbox(slide, Emu(x + boxw), y, gap, boxh, [("→", 24, ACCENT, True)],
                    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        x = Emu(x + boxw + gap)


def table(slide, x, y, w, h, data):
    rows, cols = len(data), len(data[0])
    g = slide.shapes.add_table(rows, cols, x, y, w, h).table
    for ri, row in enumerate(data):
        for ci, val in enumerate(row):
            cell = g.cell(ri, ci); cell.text = str(val)
            para = cell.text_frame.paragraphs[0]
            run = para.runs[0]; run.font.name = "Calibri"; run.font.size = Pt(15)
            if ri == 0:
                run.font.bold = True; run.font.color.rgb = WHITE
                cell.fill.solid(); cell.fill.fore_color.rgb = ACCENT
            else:
                run.font.color.rgb = DARK
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE if ri % 2 else CARD
    return g


# ============================================================ 1 — Title
s = prs.slides.add_slide(TITLE_SLIDE)
rect(s, 0, Inches(6.9), SW, Inches(0.6), ACCENT)
rect(s, 0, Inches(6.82), SW, Pt(4), GREEN)
t = s.shapes.title
t.left, t.top, t.width, t.height = Inches(1), Inches(2.0), Inches(11.3), Inches(1.5)
t.text_frame.paragraphs[0].text = "QuizAI"
tr = t.text_frame.paragraphs[0].runs[0]
tr.font.size = Pt(60); tr.font.bold = True; tr.font.color.rgb = ACCENT_DK
t.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
sub = s.placeholders[1]
sub.left, sub.top, sub.width, sub.height = Inches(1), Inches(3.4), Inches(11.3), Inches(2.6)
stf = sub.text_frame; stf.word_wrap = True
lines = [
    ("AI-Powered Quiz Generator from PDF & PowerPoint", 24, DARK, False),
    ("Artificial Intelligence Project  ·  Topic: Generative AI", 18, ACCENT, True),
    ("Group Members:  1. Mohamed Gasem · 2. Maulid Yuswan Hidayat · 3. Mohamed Ahmed", 15, GRAY, False),
    ("github.com/Mirror71/Quiz-Ai-", 13, GRAY, False),
]
for i, (text, size, color, bold) in enumerate(lines):
    p = stf.paragraphs[0] if i == 0 else stf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER; p.space_after = Pt(8)
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold; r.font.name = "Calibri"

# ============================================================ 2 — Problem & Solution
s = content_slide("The Problem & Our Solution")
textbox(s, Inches(0.6), Inches(1.6), Inches(5.7), Inches(0.5), [("❗ The Problem", 22, ACCENT_DK, True)])
bullets(s, Inches(0.6), Inches(2.2), Inches(5.7), Inches(2.6), [
    "Creating quiz questions from study material is slow and tedious.",
    "Students need fast self-assessment to test understanding.",
    "Good options + explanations need effort and expertise.",
], size=18)
# center divider
rect(s, Inches(6.62), Inches(1.7), Pt(2), Inches(3.0), RGBColor(0xD0, 0xD7, 0xE2))
textbox(s, Inches(7.0), Inches(1.6), Inches(5.7), Inches(0.5), [("✅ Our Solution: QuizAI", 22, ACCENT_DK, True)])
bullets(s, Inches(7.0), Inches(2.2), Inches(5.7), Inches(2.6), [
    "Upload any PDF or PowerPoint document.",
    "Gemini AI reads it and writes a 10-question quiz.",
    "Take it interactively: feedback, explanations & a score.",
], size=18)
# bottom transformation banner (fills empty space)
rect(s, Inches(1.3), Inches(5.4), Inches(10.7), Inches(1.2), CARD)
textbox(s, Inches(1.3), Inches(5.4), Inches(10.7), Inches(1.2),
        [("📄  Upload Document     →     🤖  Gemini AI Reads It     →     ❓  10-Question Quiz", 20, ACCENT_DK, True)],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ============================================================ 3 — How it works
s = content_slide("How It Works")
steps = [
    ("1", "Upload", "User drops a PDF / PPTX (max 10 MB)"),
    ("2", "Extract", "Server extracts text (pdf-parse / slide XML)"),
    ("3", "Generate", "Gemini returns a 10-question quiz as JSON"),
    ("4", "Validate", "JSON parsed, repaired & schema-checked"),
    ("5", "Play", "Interactive quiz → feedback → score"),
]
n = len(steps); gap = Inches(0.3)
cardw = Emu(int((SW - Inches(1.2) - gap * (n - 1)) / n))
x = Inches(0.6); y = Inches(2.4); ch = Inches(2.5)
for num, ttl, desc in steps:
    rect(s, x, y, cardw, ch, CARD)
    rect(s, x, y, cardw, Inches(0.7), ACCENT)
    textbox(s, x, y, cardw, Inches(0.7), [(num, 26, WHITE, True)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, x, y + Inches(0.85), cardw, Inches(0.5), [(ttl, 17, ACCENT_DK, True)], align=PP_ALIGN.CENTER)
    textbox(s, Emu(x + Inches(0.08)), y + Inches(1.35), Emu(cardw - Inches(0.16)), Inches(1.1),
            [(desc, 12, GRAY, False)], align=PP_ALIGN.CENTER)
    x = Emu(x + cardw + gap)
textbox(s, Inches(0.6), Inches(5.5), Inches(12.1), Inches(0.8),
        [("Browser (React)  →  Express API (Node.js)  →  Google Gemini API", 18, ACCENT_DK, True)],
        align=PP_ALIGN.CENTER)

# ============================================================ 4 — Tech & Features
s = content_slide("Technology & Features")
textbox(s, Inches(0.6), Inches(1.6), Inches(5.9), Inches(0.5), [("🛠️ Tech Stack", 22, ACCENT_DK, True)])
bullets(s, Inches(0.6), Inches(2.2), Inches(5.9), Inches(2.6), [
    "Frontend: React + Vite + Tailwind CSS",
    "Backend: Node.js + Express",
    "AI Model: Google Gemini (gemini-2.5-flash)",
    "Parsing: pdf-parse + JSZip (PPTX)",
], size=18)
textbox(s, Inches(6.9), Inches(1.6), Inches(5.9), Inches(0.5), [("✨ Key Features", 22, ACCENT_DK, True)])
bullets(s, Inches(6.9), Inches(2.2), Inches(5.9), Inches(2.6), [
    "Drag-and-drop upload with validation.",
    "Instant green / red answer feedback.",
    "Score screen with per-question review.",
    "Robust: auto-retry + model fallback.",
], size=18)
# mini architecture diagram (fills bottom)
textbox(s, Inches(0.6), Inches(4.9), Inches(12), Inches(0.4), [("System Architecture", 16, GRAY, True)])
flow(s, Inches(5.4), ["React Frontend\n(browser)", "Express API\n(Node.js)", "Google Gemini API\n(AI model)"], boxh=Inches(1.0))

# ============================================================ 5 — Dataset & Method
s = content_slide("Dataset & Evaluation Method")
textbox(s, Inches(0.6), Inches(1.6), Inches(5.9), Inches(0.5), [("📚 Dataset", 22, ACCENT_DK, True)])
bullets(s, Inches(0.6), Inches(2.2), Inches(5.9), Inches(2.3), [
    "10 fact-verified educational documents.",
    "6 domains (Biology, Physics, History…).",
    "Avg ~1,750 characters each; sources cited.",
], size=18)
textbox(s, Inches(6.9), Inches(1.6), Inches(5.9), Inches(0.5), [("🔬 How We Evaluated", 22, ACCENT_DK, True)])
bullets(s, Inches(6.9), Inches(2.2), Inches(5.9), Inches(2.3), [
    "Generated quizzes across 3 Gemini models.",
    "Measured success, latency, JSON validity.",
    "Checked answer-position balance (bias).",
], size=18)
# dataset summary table (fills bottom)
textbox(s, Inches(0.6), Inches(4.5), Inches(12), Inches(0.4), [("Dataset Summary", 16, GRAY, True)])
table(s, Inches(0.6), Inches(4.95), Inches(12.1), Inches(1.6), [
    ["Documents", "Domains", "Avg length", "Models tested", "Total quiz runs"],
    ["10", "6", "~1,750 chars", "3", "30"],
])

# ============================================================ 6 — Results
s = content_slide("Results")
bullets(s, Inches(0.5), Inches(1.9), Inches(4.4), Inches(4.5), [
    "gemini-2.5-flash: 100% success.",
    "Exactly 10 questions each.",
    "100% valid JSON (first try).",
    "~12 s average per quiz.",
    "Reliable for real-time use.",
], size=17)
img = FIG / "03_success_latency.png"
if img.exists():
    s.shapes.add_picture(str(img), Inches(5.0), Inches(2.3), width=Inches(7.9))

# ============================================================ 7 — Key Finding
s = content_slide("Key Finding: Answer-Position Bias")
bullets(s, Inches(0.5), Inches(1.9), Inches(4.4), Inches(4.5), [
    "Model favored some answer slots.",
    "Chose 'C' 56% of the time…",
    "…and never 'A' (0%).",
    "Known LLM bias (Zhao 2021).",
    "Fix: shuffle options per question.",
], size=17)
img = FIG / "04_answer_bias.png"
if img.exists():
    s.shapes.add_picture(str(img), Inches(5.0), Inches(1.9), width=Inches(7.9))

# ============================================================ 8 — App Demo
s = content_slide("Application Demo")
shots = [
    ("Upload Screen", ["Drag & drop a", "PDF / PPTX file", "→ Generate Quiz"]),
    ("Quiz Screen", ["Question 3 of 10", "[A] [B] [C] [D]", "instant feedback"]),
    ("Score Screen", ["You scored 8/10", "80%  🎉", "review answers"]),
]
n = len(shots); gap = Inches(0.6)
cw = Emu(int((SW - Inches(1.2) - gap * (n - 1)) / n))
x = Inches(0.6); y = Inches(1.9); chh = Inches(3.9)
for ttl, body in shots:
    rect(s, x, y, cw, chh, CARD)
    rect(s, x, y, cw, Inches(0.65), ACCENT)
    textbox(s, x, y, cw, Inches(0.65), [(ttl, 18, WHITE, True)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, Emu(x + Inches(0.2)), y + Inches(1.1), Emu(cw - Inches(0.4)), Inches(2.5),
            [(line, 16, GRAY, False) for line in body], align=PP_ALIGN.CENTER)
    x = Emu(x + cw + gap)
textbox(s, Inches(0.6), Inches(6.1), Inches(12.1), Inches(0.6),
        [("Live demo at http://localhost:5000  ·  (replace these boxes with real screenshots)", 14, GRAY, True)],
        align=PP_ALIGN.CENTER)

# ============================================================ 9 — Conclusion
s = content_slide("Conclusion")
bullets(s, Inches(0.7), Inches(1.7), Inches(11.9), Inches(2.5), [
    "Built a complete Generative AI app: documents → interactive quizzes.",
    "gemini-2.5-flash generates reliable, well-structured quizzes (100% success).",
    "Evaluation revealed a measurable, fixable answer-position bias.",
    "Deployed and running locally; full code on GitHub.",
], size=19)
textbox(s, Inches(0.7), Inches(4.5), Inches(11.9), Inches(0.5), [("🚀 Future Work", 20, ACCENT_DK, True)])
flow(s, Inches(5.1), ["Shuffle answer\noptions (fix bias)", "OCR for scanned\ndocuments", "Configurable\ndifficulty & length", "Public cloud\ndeployment"], boxh=Inches(1.0))
rect(s, 0, Inches(6.9), SW, Inches(0.6), ACCENT)
textbox(s, Inches(1), Inches(6.92), Inches(11.3), Inches(0.5),
        [("github.com/Mirror71/Quiz-Ai-      ·      Thank you!", 16, WHITE, True)],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ============================================================ 10 — References
s = content_slide("References (Daftar Pustaka)")
refs = [
    "Brown, T. B., et al. (2020). Language models are few-shot learners. NeurIPS, 33.",
    "Gemini Team, Google. (2023). Gemini: A family of highly capable multimodal models. arXiv:2312.11805.",
    "Kasneci, E., et al. (2023). ChatGPT for good? On opportunities and challenges of large language models for education. Learning and Individual Differences, 103.",
    "Kurdi, G., et al. (2020). A systematic review of automatic question generation for educational purposes. IJAIED, 30(1).",
    "Zhao, Z., et al. (2021). Calibrate before use: Improving few-shot performance of language models. ICML, 139.",
]
tb = s.shapes.add_textbox(Inches(0.7), Inches(1.8), Inches(12.0), Inches(5.0))
tf = tb.text_frame; tf.word_wrap = True
for i, r in enumerate(refs):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(16)
    run = p.add_run(); run.text = f"{i+1}.  {r}"
    run.font.size = Pt(16); run.font.color.rgb = DARK; run.font.name = "Calibri"

# ============================================================ Speaker notes
NOTES = [
    "Introduce yourselves and the project: QuizAI, an AI project under the "
    "Generative AI topic. One sentence: 'QuizAI turns any PDF or PowerPoint into "
    "an interactive multiple-choice quiz using Google Gemini.' Group of 3, code "
    "on GitHub.",
    "Motivation: making quiz questions by hand is slow; students need quick "
    "self-assessment. Solution: upload a document, the AI writes a 10-question "
    "quiz with answers and explanations, taken interactively. Point at the bottom "
    "banner showing the transformation.",
    "Walk through the five steps left to right. Emphasize step 3 (Gemini returns "
    "strict JSON) and step 4 (we validate/repair it). Bottom line shows the "
    "architecture: React, Express, Gemini.",
    "Highlight React, Node/Express, and Gemini. Point to features that show "
    "polish: instant feedback, score/review, robustness. Use the mini diagram to "
    "show the three-tier architecture.",
    "Explain we tested scientifically: 10 fact-checked docs across 6 subjects, "
    "quizzes generated across 3 models, objective metrics measured. Walk the "
    "audience through the summary table.",
    "Headline numbers for gemini-2.5-flash: 100% success, exactly 10 questions, "
    "100% valid JSON first try, ~12 seconds. Point at the chart comparing success "
    "and latency.",
    "Our most interesting result: position bias — chose C 56% and never A. Matches "
    "a known research finding (Zhao et al., 2021). Real quality issue; simple fix "
    "is shuffling options.",
    "Show the app flow: upload screen, quiz screen with feedback, score screen. If "
    "possible, do a LIVE demo here at localhost:5000, or replace the boxes with "
    "real screenshots beforehand.",
    "Wrap up: we built, evaluated, and deployed a complete Generative AI app. "
    "Mention future work (shuffle options, OCR, difficulty, cloud). Thank the "
    "audience and invite questions.",
    "Briefly note the key academic sources that informed the project, especially "
    "Zhao et al. (2021) for the bias finding and the Gemini technical report.",
]
for slide, note in zip(prs.slides, NOTES):
    slide.notes_slide.notes_text_frame.text = note

prs.save(OUT)
print("Wrote", OUT, "-", len(list(prs.slides)), "slides")
